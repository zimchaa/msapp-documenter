#!/usr/bin/env python3
"""Create assets/default_template.pptx: python-pptx's built-in Office theme, widened to 16:9.

Any company .pptx can replace it (see references/pptx_template.md). The renderer only relies
on layout names ("Title Slide", "Section Header", "Title Only") and on theme colours/fonts.
"""
import os
import sys

from pptx import Presentation
from pptx.util import Emu, Inches


def main(out):
    prs = Presentation()
    old_w = prs.slide_width
    new_w = Inches(13.333)
    scale = new_w / old_w
    prs.slide_width = new_w
    prs.slide_height = Inches(7.5)
    masters = list(prs.slide_masters)
    for m in masters:
        for shape_tree in [m.shapes] + [l.shapes for l in m.slide_layouts]:
            for sh in shape_tree:
                # Read (possibly inherited) geometry first and write all four values: setting only
                # `left` on an inheriting placeholder would create an xfrm with zero top/height.
                l, t, w, h = sh.left, sh.top, sh.width, sh.height
                if None in (l, t, w, h):
                    continue
                sh.left, sh.top, sh.width, sh.height = Emu(int(l * scale)), Emu(t), Emu(int(w * scale)), Emu(h)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    prs.save(out)
    print("wrote", out)


if __name__ == "__main__":
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "assets", "default_template.pptx"))
